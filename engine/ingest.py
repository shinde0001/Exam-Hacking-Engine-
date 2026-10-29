"""
Exam-Hacking Engine — Ingestion Pipeline
Handles ingestion of PDF, DOCX, TXT, and image documents.
Features:
- Digital text extraction via pdfplumber / PyMuPDF / pdftotext
- OCR fallback (pytesseract) when extracted text < OCR_CHAR_THRESHOLD per page
- Incremental processing tracking (status: pending -> ingesting -> ingested / failed)
- SHA-256 deduplication to skip already processed files
- Automatic year detection from filenames
"""

import hashlib
import os
import re
import shutil
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from typing import Any, List, Optional, Tuple

import config
from engine import db

# Optional imports handled gracefully
try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    import pymupdf as fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz
        HAS_PYMUPDF = True
    except ImportError:
        HAS_PYMUPDF = False

try:
    import docx  # python-docx
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False

try:
    import pytesseract
    from PIL import Image
    HAS_PYTESSERACT = True
except ImportError:
    HAS_PYTESSERACT = False


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    char_count: int
    method: str
    ocr_applied: bool


@dataclass
class IngestionResult:
    document_id: int
    filename: str
    file_hash: str
    status: str
    total_pages: int
    pages: List[ExtractedPage]
    skipped: bool = False
    message: str = ""


def compute_file_hash(filepath: str) -> str:
    """Calculate SHA-256 hash of a file for change tracking and deduplication."""
    hasher = hashlib.sha256()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def detect_year_from_name(filename: str) -> Optional[int]:
    """
    Auto-detect a 4-digit exam year (1900-2099) from filename.
    Handles underscores, hyphens, brackets, dots, e.g.:
    - 'Math_2022_Final.pdf' -> 2022
    - 'Exam-Paper(2023).docx' -> 2023
    """
    matches = re.findall(r"(?<!\d)(19\d{2}|20\d{2})(?!\d)", filename)
    if matches:
        return int(matches[-1])  # Use most recent or trailing year
    return None


def is_tesseract_available() -> bool:
    """Check if the system tesseract binary is installed and executable."""
    return shutil.which("tesseract") is not None and HAS_PYTESSERACT


def run_ocr_on_image(image: Any) -> str:
    """Run pytesseract OCR on a PIL Image object with configured parameters."""
    if not is_tesseract_available():
        return ""
    try:
        text = pytesseract.image_to_string(image, config=config.TESSERACT_CONFIG)
        return text.strip()
    except Exception as e:
        print(f"[Ingest OCR] OCR failed: {e}")
        return ""


def extract_text_from_pdf(filepath: str) -> List[ExtractedPage]:
    """
    Extract text per page from PDF using pdfplumber, with PyMuPDF and pdftotext fallbacks.
    Triggers OCR if extracted characters on a page fall below OCR_CHAR_THRESHOLD.
    """
    pages: List[ExtractedPage] = []

    # Strategy 1: pdfplumber (Primary as per plan)
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(filepath) as pdf:
                for idx, page in enumerate(pdf.pages, start=1):
                    digital_text = (page.extract_text() or "").strip()
                    ocr_applied = False
                    method = "pdfplumber"

                    if len(digital_text) < config.OCR_CHAR_THRESHOLD:
                        if is_tesseract_available():
                            try:
                                pil_image = page.to_image(resolution=200).original
                                ocr_text = run_ocr_on_image(pil_image)
                                if len(ocr_text) > len(digital_text):
                                    digital_text = ocr_text
                                    ocr_applied = True
                                    method = "pdfplumber_ocr"
                                if hasattr(pil_image, "close"):
                                    pil_image.close()
                                del pil_image
                            except Exception as e:
                                print(f"[Ingest PDF] Page {idx} OCR fallback error: {e}")

                    pages.append(
                        ExtractedPage(
                            page_number=idx,
                            text=digital_text,
                            char_count=len(digital_text),
                            method=method,
                            ocr_applied=ocr_applied,
                        )
                    )
            if pages:
                return pages
        except Exception as e:
            print(f"[Ingest PDF] pdfplumber failed: {e}. Trying fallback.")

    # Strategy 2: PyMuPDF (fitz)
    if HAS_PYMUPDF:
        try:
            doc = fitz.open(filepath)
            for idx, page in enumerate(doc, start=1):
                digital_text = page.get_text().strip()
                ocr_applied = False
                method = "pymupdf"

                if len(digital_text) < config.OCR_CHAR_THRESHOLD:
                    if is_tesseract_available():
                        try:
                            pix = page.get_pixmap(dpi=200)
                            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                            ocr_text = run_ocr_on_image(img)
                            if len(ocr_text) > len(digital_text):
                                digital_text = ocr_text
                                ocr_applied = True
                                method = "pymupdf_ocr"
                            img.close()
                            del pix, img
                        except Exception as e:
                            print(f"[Ingest PyMuPDF] Page {idx} OCR error: {e}")

                pages.append(
                    ExtractedPage(
                        page_number=idx,
                        text=digital_text,
                        char_count=len(digital_text),
                        method=method,
                        ocr_applied=ocr_applied,
                    )
                )
            doc.close()
            if pages:
                return pages
        except Exception as e:
            print(f"[Ingest PDF] PyMuPDF failed: {e}. Trying system pdftotext.")

    # Strategy 3: pdftotext CLI (Standard Linux tool)
    pdftotext_path = shutil.which("pdftotext")
    if pdftotext_path:
        try:
            # Extract entire document separated by form feeds (\x0c)
            result = subprocess.run(
                [pdftotext_path, "-layout", filepath, "-"],
                capture_output=True,
                text=True,
                check=True,
            )
            raw_pages = result.stdout.split("\x0c")
            # Last split is often empty after final form-feed
            if raw_pages and not raw_pages[-1].strip():
                raw_pages = raw_pages[:-1]

            for idx, page_str in enumerate(raw_pages, start=1):
                clean_text = page_str.strip()
                pages.append(
                    ExtractedPage(
                        page_number=idx,
                        text=clean_text,
                        char_count=len(clean_text),
                        method="pdftotext_cli",
                        ocr_applied=False,
                    )
                )
            if pages:
                return pages
        except Exception as e:
            print(f"[Ingest PDF] pdftotext CLI error: {e}")

    # If all PDF extractors failed
    if not pages:
        raise RuntimeError(f"Unable to extract text from PDF '{filepath}'. No suitable extractor succeeded.")

    return pages


def extract_text_from_docx(filepath: str) -> List[ExtractedPage]:
    """
    Extract text from DOCX file.
    Uses python-docx if installed, otherwise parses word/document.xml directly via zipfile.
    """
    paragraphs: List[str] = []

    if HAS_PYTHON_DOCX:
        try:
            doc = docx.Document(filepath)
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
        except Exception as e:
            print(f"[Ingest DOCX] python-docx parsing error: {e}. Falling back to XML parser.")
            paragraphs = []

    # Standalone XML Fallback
    if not paragraphs:
        try:
            with zipfile.ZipFile(filepath) as z:
                xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            # Find all paragraph tags <w:p>
            for p in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                texts = [
                    t.text for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if t.text
                ]
                if texts:
                    paragraphs.append("".join(texts).strip())
        except Exception as e:
            raise RuntimeError(f"Failed to parse DOCX '{filepath}': {e}")

    full_text = "\n\n".join(paragraphs).strip()
    return [
        ExtractedPage(
            page_number=1,
            text=full_text,
            char_count=len(full_text),
            method="python_docx" if HAS_PYTHON_DOCX and paragraphs else "docx_xml",
            ocr_applied=False,
        )
    ]


def extract_text_from_txt(filepath: str) -> List[ExtractedPage]:
    """
    Extract text from plain text or markdown file.
    Splits by form feed (\\x0c) if present, otherwise returns as page 1.
    """
    content = ""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(filepath, "r", encoding=encoding) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue

    if not content:
        raise ValueError(f"File '{filepath}' is empty or could not be decoded.")

    raw_pages = content.split("\x0c")
    if raw_pages and not raw_pages[-1].strip():
        raw_pages = raw_pages[:-1]

    if not raw_pages:
        raw_pages = [content]

    pages = []
    for idx, page_str in enumerate(raw_pages, start=1):
        clean = page_str.strip()
        pages.append(
            ExtractedPage(
                page_number=idx,
                text=clean,
                char_count=len(clean),
                method="text_read",
                ocr_applied=False,
            )
        )
    return pages


def extract_text_from_image(filepath: str) -> List[ExtractedPage]:
    """
    Extract text from an image file using OCR.
    """
    if not is_tesseract_available():
        raise RuntimeError("OCR is required for images, but Tesseract/pytesseract is not available.")

    try:
        with Image.open(filepath) as img:
            text = run_ocr_on_image(img)
            return [
                ExtractedPage(
                    page_number=1,
                    text=text,
                    char_count=len(text),
                    method="image_ocr",
                    ocr_applied=True,
                )
            ]
    except Exception as e:
        raise RuntimeError(f"Failed to process image '{filepath}': {e}")


def extract_document_content(filepath: str) -> Tuple[str, List[ExtractedPage]]:
    """
    Route file by extension to the appropriate extractor.
    Returns (file_type, pages).
    """
    ext = os.path.splitext(filepath)[1].lower().lstrip(".")
    if ext not in config.ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: .{ext}")

    if ext == "pdf":
        return "pdf", extract_text_from_pdf(filepath)
    elif ext in ("docx", "doc"):
        return "docx", extract_text_from_docx(filepath)
    elif ext in ("txt", "md"):
        return "txt", extract_text_from_txt(filepath)
    elif ext in ("png", "jpg", "jpeg", "tiff", "bmp"):
        return "image", extract_text_from_image(filepath)
    else:
        raise ValueError(f"No extractor implemented for .{ext}")


def ingest_file(
    filepath: str,
    year: Optional[int] = None,
    subject: Optional[str] = None,
    force_reingest: bool = False,
    db_path: Optional[str] = None,
) -> IngestionResult:
    """
    Ingest a single document file into the Exam-Hacking Engine.
    Implements incremental processing:
    - Calculates SHA-256 hash.
    - If document already exists and is in 'ingested'/'extracted'/'indexed' status and force_reingest is False, skips.
    - If status was 'failed', or force_reingest is True, re-processes.
    - Records metadata, raw text per page, character counts, and extraction methods in SQLite.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    filename = os.path.basename(filepath)
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext not in config.ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: .{ext}")

    file_size = os.path.getsize(filepath)
    file_hash = compute_file_hash(filepath)

    # Auto-detect year if not supplied
    if year is None:
        year = detect_year_from_name(filename)

    # Incremental Processing Check
    existing_doc = db.get_document_by_hash(file_hash, db_path=db_path)
    if existing_doc and not force_reingest:
        status = existing_doc["status"]
        if status in ("ingested", "extracted", "indexed"):
            pages_records = db.get_document_pages(existing_doc["id"], db_path=db_path)
            extracted_pages = [
                ExtractedPage(
                    page_number=p["page_number"],
                    text=p["raw_text"],
                    char_count=p["char_count"],
                    method=p["extraction_method"],
                    ocr_applied=bool(p["ocr_applied"]),
                )
                for p in pages_records
            ]
            return IngestionResult(
                document_id=existing_doc["id"],
                filename=existing_doc["filename"],
                file_hash=file_hash,
                status=status,
                total_pages=existing_doc["total_pages"],
                pages=extracted_pages,
                skipped=True,
                message=f"Document already ingested (status: {status}). Skipped.",
            )

    # Determine or insert document ID
    if existing_doc:
        doc_id = existing_doc["id"]
        db.update_document_status(doc_id, "ingesting", error_message=None, db_path=db_path)
    else:
        ext = os.path.splitext(filename)[1].lower().lstrip(".")
        doc_id = db.insert_document(
            filename=filename,
            filepath=filepath,
            file_hash=file_hash,
            file_type=ext,
            file_size=file_size,
            year=year,
            subject=subject,
            db_path=db_path,
        )
        db.update_document_status(doc_id, "ingesting", db_path=db_path)

    try:
        file_type, pages = extract_document_content(filepath)

        # Store pages in database
        for page in pages:
            db.insert_document_page(
                doc_id=doc_id,
                page_number=page.page_number,
                raw_text=page.text,
                char_count=page.char_count,
                extraction_method=page.method,
                ocr_applied=page.ocr_applied,
                db_path=db_path,
            )

        # Mark document as successfully ingested
        db.update_document_status(doc_id, "ingested", total_pages=len(pages), db_path=db_path)

        return IngestionResult(
            document_id=doc_id,
            filename=filename,
            file_hash=file_hash,
            status="ingested",
            total_pages=len(pages),
            pages=pages,
            skipped=False,
            message=f"Successfully ingested {len(pages)} page(s).",
        )

    except Exception as e:
        error_msg = str(e)
        db.update_document_status(doc_id, "failed", error_message=error_msg, db_path=db_path)
        raise RuntimeError(f"Ingestion failed for document ID {doc_id} ('{filename}'): {error_msg}") from e
